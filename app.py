import io
import json
import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import mysql.connector
import pandas as pd
import streamlit as st
from langchain_community.document_loaders import PyPDFLoader
from langchain_community.embeddings import HuggingFaceEmbeddings
from langchain_community.vectorstores import Chroma
from langchain_core.output_parsers import StrOutputParser
from langchain_core.prompts import ChatPromptTemplate
from langchain_core.runnables import RunnablePassthrough
from langchain_text_splitters import RecursiveCharacterTextSplitter

# Opsional untuk LLM Cloud agar bisa diakses di Streamlit Cloud
from langchain_groq import ChatGroq
from langchain_community.llms import Ollama

# 1. Konfigurasi Halaman (Wide Mode)
st.set_page_config(
    page_title="EduBoard AI: UKRIDA Local RAG Ground Truth",
    page_icon="🎬",
    layout="wide",
)

# Custom CSS dengan Tema Warna Khas UKRIDA & Styling Tabel
st.markdown(
    """
    <style>
        .main-header { 
            font-size: 26px; 
            font-weight: bold; 
            color: #003366; 
        }
        .sub-text { 
            color: #2B6CB0; 
            font-size: 14px; 
            margin-bottom: 20px; 
        }
        .card-box { 
            background-color: #F0F4F8; 
            padding: 15px; 
            border-radius: 8px; 
            border-left: 4px solid #003366; 
            border-top: 1px solid #E2E8F0; 
            border-right: 1px solid #E2E8F0; 
            border-bottom: 1px solid #E2E8F0; 
            margin-bottom: 10px; 
        }
        .stButton>button {
            background-color: #003366;
            color: white;
            border-radius: 6px;
            font-weight: 600;
        }
        .stButton>button:hover {
            background-color: #D69E2E;
            color: #003366;
        }
    </style>
""",
    unsafe_allow_html=True,
)


# --- FUNGSI DATABASE MYSQL (KONEKSI RAILWAY) ---
def get_db_connection():
  if "mysql" in st.secrets:
    db_config = st.secrets["mysql"]
    host = db_config.get("host", "localhost")
    user = db_config.get("user", "root")
    password = db_config.get("password", "")
    database = db_config.get("database", "railway")
    port = int(db_config.get("port", 3306))
  else:
    host = "localhost"
    user = "root"
    password = ""
    database = "railway"
    port = 3306

  return mysql.connector.connect(
      host=host,
      user=user,
      password=password,
      database=database,
      port=port,
      connect_timeout=10,
  )


def check_db_connection():
  try:
    conn = get_db_connection()
    conn.close()
    return True
  except mysql.connector.Error:
    return False


def init_db_tables():
  try:
    if check_db_connection():
      conn = get_db_connection()
      cursor = conn.cursor()
      cursor.execute("""
                CREATE TABLE IF NOT EXISTS users (
                    id INT AUTO_INCREMENT PRIMARY KEY,
                    username VARCHAR(100) UNIQUE NOT NULL,
                    password VARCHAR(255) NOT NULL,
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
            """)
      cursor.execute("""
                CREATE TABLE IF NOT EXISTS storyboards (
                    id INT AUTO_INCREMENT PRIMARY KEY,
                    username VARCHAR(100),
                    program_studi VARCHAR(255),
                    nama_mata_kuliah VARCHAR(255),
                    project_name VARCHAR(255),
                    learning_objectives TEXT,
                    target_audience VARCHAR(100),
                    visual_style VARCHAR(100),
                    rps_filename VARCHAR(255) DEFAULT '-',
                    created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                )
            """)
      cursor.execute("""
                CREATE TABLE IF NOT EXISTS storyboard_scenes (
                    id INT AUTO_INCREMENT PRIMARY KEY,
                    storyboard_id INT,
                    judul_scene VARCHAR(255),
                    visualisasi TEXT,
                    instruksi_visual TEXT,
                    animasi TEXT,
                    on_screen_text TEXT,
                    voice_over_text TEXT,
                    backsound VARCHAR(100),
                    durasi VARCHAR(50),
                    saran_info TEXT,
                    FOREIGN KEY (storyboard_id) REFERENCES storyboards(id) ON DELETE CASCADE
                )
            """)
      conn.commit()
      cursor.close()
      conn.close()
  except Exception:
    pass


init_db_tables()


def parse_llm_json_response(response_text):
  try:
    text = response_text.strip()
    if "```" in text:
      text = re.sub(r"^```[a-zA-Z]*\n?", "", text)
      text = re.sub(r"\n?```$", "", text)
      text = text.strip()
    try:
      return json.loads(text)
    except Exception:
      pass
    match = re.search(r"\[\s*\{.*\}\s*\]", text, re.DOTALL)
    if match:
      return json.loads(match.group(0))
    return None
  except Exception:
    return None


def generate_pptx(df, program_studi, nama_mata_kuliah, project_title):
  prs = Presentation()
  for index, row in df.iterrows():
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    table_shape = slide.shapes.add_table(
        5, 4, Inches(0.8), Inches(0.8), Inches(11.7), Inches(6.0)
    )
    table = table_shape.table
    table.columns[0].width, table.columns[1].width = Inches(3.2), Inches(3.2)
    table.columns[2].width, table.columns[3].width = Inches(3.0), Inches(2.3)

    table.cell(0, 0).text = f"Program Studi\n{program_studi}"
    table.cell(0, 1).text = ":"
    cell_jdl = table.cell(0, 2)
    cell_jdl.merge(table.cell(0, 3))
    cell_jdl.text = f"Judul Scene:\n{row.get('Judul Scene', '')}"

    table.cell(1, 0).text = f"Nama Mata Kuliah\n{nama_mata_kuliah}"
    table.cell(1, 1).text = ":"

    table.cell(2, 0).text = "Visualisasi"
    table.cell(2, 2).text = "Instruksi untuk Visualisasi"
    table.cell(2, 3).text = "Animasi"
    cell_vis = table.cell(2, 0)
    cell_vis.merge(table.cell(3, 1))
    cell_vis.text = f"Visualisasi:\n{row.get('Visualisasi', '')}"
    table.cell(2, 2).text = str(row.get("Instruksi untuk Visual", ""))
    table.cell(2, 3).text = str(row.get("Animasi", ""))

    table.cell(4, 0).text = f"On-Screen Text:\n{row.get('On-Screen Text', '')}"
    table.cell(4, 1).text = f"Voice Over Text:\n{row.get('Voice Over Text', '')}"
    table.cell(4, 2).text = f"Backsound:\n{row.get('Backsound', '')}"
    table.cell(4, 3).text = f"Durasi:\n{row.get('Durasi', '')}"

    if row.get("Saran dan Info"):
      saran_box = slide.shapes.add_textbox(
          Inches(0.8), Inches(6.9), Inches(11.7), Inches(0.5)
      )
      tf = saran_box.text_frame
      tf.word_wrap = True
      tf.paragraphs[0].text = f"Saran / Info: {row.get('Saran dan Info', '')}"
      tf.paragraphs[0].font.size = Pt(11)

    for r_idx, r_cells in enumerate(table.rows):
      for c_idx, cell in enumerate(r_cells.cells):
        cell.fill.solid()
        if r_idx in [0, 1, 2] and c_idx in [0, 2, 3]:
          cell.fill.fore_color.rgb = RGBColor(0, 51, 102)
          for p in cell.text_frame.paragraphs:
            p.font.bold = True
            p.font.color.rgb = RGBColor(255, 255, 255)
            p.font.size = Pt(12)
        else:
          cell.fill.fore_color.rgb = RGBColor(240, 244, 248)
          for p in cell.text_frame.paragraphs:
            p.font.color.rgb = RGBColor(51, 51, 51)
            p.font.size = Pt(11)

  ppt_stream = io.BytesIO()
  prs.save(ppt_stream)
  ppt_stream.seek(0)
  return ppt_stream


# --- MANAJEMEN SESSION AUTH ---
if "logged_in" not in st.session_state:
  st.session_state.logged_in = False
if "username" not in st.session_state:
  st.session_state.username = ""
if "is_admin" not in st.session_state:
  st.session_state.is_admin = False

# ==========================================
# KONDISI 1: JIKA BELUM LOGIN
# ==========================================
if not st.session_state.logged_in:
  st.markdown(
      '<div style="text-align: center; margin-top: 50px;" class="main-header">'
      " UKRIDA EduBoard Ground Truth AI 🎓🎬</div>",
      unsafe_allow_html=True,
  )
  st.markdown(
      '<div style="text-align: center;" class="sub-text">Automated'
      " Storyboard Generator based on Ground Truth & Local RAG</div>",
      unsafe_allow_html=True,
  )

  col1, col2, col3 = st.columns([1, 1.2, 1])
  with col2:
    if not check_db_connection():
      st.error(
          "⚠️ **Database Railway Tidak Terhubung!** Periksa parameter host dan"
          " port publik pada st.secrets."
      )

    tab_login, tab_register = st.tabs(["🔐 Login", "📝 Register"])

    with tab_login:
      st.subheader("Masuk ke Akun Anda")
      login_user = st.text_input("Username", key="login_user")
      login_pass = st.text_input(
          "Password", type="password", key="login_pass"
      )

      if st.button("Login", use_container_width=True, type="primary"):
        if not login_user or not login_pass:
          st.warning("Isi username dan password terlebih dahulu.")
        else:
          try:
            conn = get_db_connection()
            cursor = conn.cursor(dictionary=True)
            cursor.execute(
                "SELECT * FROM users WHERE username = %s AND password = %s",
                (login_user, login_pass),
            )
            user = cursor.fetchone()
            cursor.close()
            conn.close()

            if user:
              st.session_state.logged_in = True
              st.session_state.username = user["username"]
              st.session_state.is_admin = (
                  user["username"].lower() == "admin"
              )
              st.session_state.storyboard_df = pd.DataFrame(columns=[
                  "Judul Scene",
                  "Visualisasi",
                  "Instruksi untuk Visual",
                  "Animasi",
                  "On-Screen Text",
                  "Voice Over Text",
                  "Backsound",
                  "Durasi",
                  "Saran dan Info",
              ])
              st.session_state.uploaded_rps_name = "-"
              st.session_state.retriever = None
              st.success("Login berhasil!")
              st.rerun()
            else:
              st.error("Username atau password salah.")
          except Exception as e:
            st.error(f"Terjadi kesalahan database: {e}")

    with tab_register:
      st.subheader("Buat Akun Baru")
      reg_user = st.text_input("Pilih Username", key="reg_user")
      reg_pass = st.text_input(
          "Pilih Password", type="password", key="reg_pass"
      )
      reg_pass_confirm = st.text_input(
          "Konfirmasi Password", type="password", key="reg_pass_confirm"
      )

      if st.button("Daftar Akun", use_container_width=True):
        if not reg_user or not reg_pass:
          st.warning("Username dan password tidak boleh kosong.")
        elif reg_pass != reg_pass_confirm:
          st.error("Konfirmasi password tidak cocok.")
        else:
          try:
            conn = get_db_connection()
            cursor = conn.cursor()
            cursor.execute(
                "SELECT * FROM users WHERE username = %s", (reg_user,)
            )
            if cursor.fetchone():
              st.error("Username sudah terpakai.")
            else:
              cursor.execute(
                  "INSERT INTO users (username, password) VALUES (%s, %s)",
                  (reg_user, reg_pass),
              )
              conn.commit()
              st.success("✅ Registrasi berhasil! Silakan ke tab Login.")
            cursor.close()
            conn.close()
          except Exception as e:
            st.error(f"Gagal mendaftarkan akun: {e}")

# ==========================================
# KONDISI 2: JIKA LOGIN SEBAGAI ADMIN / USER
# ==========================================
elif st.session_state.is_admin:
  with st.sidebar:
    st.markdown("### 🛡️ Super Admin Panel")
    st.write(f"👤 Admin: **{st.session_state.username}**")
    if st.button("🚪 Logout", use_container_width=True):
      st.session_state.logged_in = False
      st.rerun()

  st.markdown(
      '<div class="main-header">🛡️ Super Admin Dashboard - UKRIDA EduBoard</div>',
      unsafe_allow_html=True,
  )
  tab_admin_users, tab_admin_projects = st.tabs(
      ["👥 Kelola Data User", "📂 Kelola Data Storyboard"]
  )

  with tab_admin_users:
    if check_db_connection():
      conn = get_db_connection()
      df_users = pd.read_sql("SELECT id, username, created_at FROM users", conn)
      conn.close()
      st.dataframe(df_users, use_container_width=True, hide_index=True)

  with tab_admin_projects:
    if check_db_connection():
      conn = get_db_connection()
      df_proj = pd.read_sql("SELECT * FROM storyboards", conn)
      conn.close()
      st.dataframe(df_proj, use_container_width=True, hide_index=True)

else:
  with st.sidebar:
    st.markdown("### 🏛️ UKRIDA EduBoard")
    st.write(f"👤 User: **{st.session_state.username}**")
    if check_db_connection():
      st.caption("🟢 MySQL Railway: Connected")
    else:
      st.caption("🔴 MySQL Railway: Disconnected")

    st.divider()
    st.markdown("### ⚙️ Konfigurasi LLM")
    llm_provider = st.selectbox(
        "Pilih Penyedia LLM",
        ["Ollama (Lokal PC)", "Groq API (Cloud / Streamlit Cloud)"],
    )

    if llm_provider == "Groq API (Cloud / Streamlit Cloud)":
      groq_api_key = st.text_input(
          "Masukkan Groq API Key", type="password", value=""
      )
      groq_model = st.selectbox(
          "Model Groq", ["llama-3.1-8b-instant", "llama-3.3-70b-versatile"]
      )
    else:
      local_model_name = st.selectbox(
          "Model Ollama", ["llama3", "mistral", "phi3"]
      )

    if st.button("🚪 Logout", use_container_width=True):
      st.session_state.logged_in = False
      st.rerun()

  st.markdown(
      '<div class="main-header">UKRIDA EduBoard AI (Ground Truth RAG'
      ' Mode)</div>',
      unsafe_allow_html=True,
  )


  @st.cache_resource
  def initialize_local_vector_store(file_path):
    loader = PyPDFLoader(file_path)
    documents = loader.load()
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000, chunk_overlap=200
    )
    docs_split = splitter.split_documents(documents)
    embeddings = HuggingFaceEmbeddings(model_name="all-MiniLM-L6-v2")
    vectorstore = Chroma.from_documents(docs_split, embeddings)
    return vectorstore.as_retriever(
        search_type="similarity", search_kwargs={"k": 4}
    )


  if "retriever" not in st.session_state:
    st.session_state.retriever = None
  if "rag_contexts" not in st.session_state:
    st.session_state.rag_contexts = []
  if "uploaded_rps_name" not in st.session_state:
    st.session_state.uploaded_rps_name = "-"

  col_left, col_middle, col_right = st.columns([1.1, 1.2, 2.7], gap="medium")

  with col_left:
    st.subheader("Project Input & Settings")
    program_studi = st.text_input("Program Studi", value="")
    nama_mata_kuliah = st.text_input("Nama Mata Kuliah", value="")
    project_name = st.text_input("Project Name", value="")
    learning_objectives = st.text_area("Learning Objectives", value="", height=120)
    target_audience = st.selectbox(
        "Target Audience",
        [
            "-- Pilih Target Audience --",
            "Mahasiswa S1/S2 (Advanced)",
            "Beginner Students",
        ],
    )
    visual_style = st.selectbox(
        "Visual Style",
        [
            "-- Pilih Visual Style --",
            "Animasi Diagram Blok & Rumus Matematis",
            "Cinematic",
        ],
    )

    uploaded_file = st.file_uploader(
        "Upload Dokumen RPS / Kurikulum (PDF)", type=["pdf"]
    )
    if st.button("Proses & Indeks Dokumen RPS"):
      if uploaded_file:
        with st.spinner("Memproses indeks dokumen..."):
          temp_path = "temp_curriculum.pdf"
          with open(temp_path, "wb") as f:
            f.write(uploaded_file.getbuffer())
          st.session_state.retriever = initialize_local_vector_store(temp_path)
          st.session_state.uploaded_rps_name = uploaded_file.name
          st.success("✅ Dokumen berhasil diindeks!")

    generate_btn = st.button(
        "Generate Storyboard", type="primary", use_container_width=True
    )

  with col_middle:
    st.subheader("RAG Content Viewer")
    if generate_btn:
      if not program_studi or not learning_objectives:
        st.error(
            "⚠️ Harap isi Program Studi dan Learning Objectives terlebih"
            " dahulu!"
        )
      elif st.session_state.retriever is None:
        st.error("⚠️ Harap upload dan proses dokumen RPS terlebih dahulu.")
      else:
        with st.spinner("🔄 Menjalankan RAG & Membangun Storyboard..."):
          try:
            relevant_docs = st.session_state.retriever.invoke(
                learning_objectives
            )
            st.session_state.rag_contexts = relevant_docs
            context_text = "\n\n".join(
                [doc.page_content for doc in relevant_docs]
            )

            if llm_provider == "Groq API (Cloud / Streamlit Cloud)":
              if not groq_api_key:
                st.error("⚠️ Masukkan Groq API Key terlebih dahulu di sidebar.")
                st.stop()
              llm = ChatGroq(groq_api_key=groq_api_key, model_name=groq_model)
            else:
              llm = Ollama(model=local_model_name, temperature=0.7)

            prompt = (
                "SYSTEM: Anda adalah API generator JSON murni. Kembalikan HANYA"
                " valid JSON array tanpa teks pengantar.\n\nBerdasarkan"
                " dokumen referensi:\n"
                f"{context_text}\n\nProgram Studi: {program_studi}\nMata Kuliah:"
                f" {nama_mata_kuliah}\nTujuan: {learning_objectives}\n\nFormat"
                " JSON:\n[\n  {\n    \"judul_scene\": \"01 - Pengenalan\","
                '\n    "visualisasi": "Ilustrasi...",\n    "instruksi_visual":'
                ' "Zoom in",\n    "animasi": "Fade in",\n    "on_screen_text":'
                ' "Teks",\n    "voice_over_text": "Narasi...",\n   '
                ' "backsound": "Calm",\n    "durasi": "00:15",\n'
                '    "saran_info": "Catatan"\n  }\n]'
            )

            response = llm.invoke(prompt)
            resp_text = (
                response.content
                if hasattr(response, "content")
                else str(response)
            )
            scenes_json = parse_llm_json_response(resp_text)

            if scenes_json and isinstance(scenes_json, list):
              formatted_scenes = []
              for item in scenes_json:
                formatted_scenes.append({
                    "Judul Scene": str(item.get("judul_scene", "Scene 01")),
                    "Visualisasi": str(item.get("visualisasi", "")),
                    "Instruksi untuk Visual": str(
                        item.get("instruksi_visual", "")
                    ),
                    "Animasi": str(item.get("animasi", "")),
                    "On-Screen Text": str(item.get("on_screen_text", "")),
                    "Voice Over Text": str(item.get("voice_over_text", "")),
                    "Backsound": str(item.get("backsound", "Audio Calm")),
                    "Durasi": str(item.get("durasi", "00:15")),
                    "Saran dan Info": str(item.get("saran_info", "")),
                })
              st.session_state.storyboard_df = pd.DataFrame(formatted_scenes)

              if check_db_connection():
                conn = get_db_connection()
                cursor = conn.cursor()
                query_main = "INSERT INTO storyboards (username, program_studi, nama_mata_kuliah, project_name, learning_objectives, target_audience, visual_style, rps_filename) VALUES (%s, %s, %s, %s, %s, %s, %s, %s)"
                cursor.execute(
                    query_main,
                    (
                        st.session_state.username,
                        program_studi,
                        nama_mata_kuliah,
                        project_name,
                        learning_objectives,
                        target_audience,
                        visual_style,
                        st.session_state.uploaded_rps_name,
                    ),
                )
                sb_id = cursor.lastrowid
                query_scene = "INSERT INTO storyboard_scenes (storyboard_id, judul_scene, visualisasi, instruksi_visual, animasi, on_screen_text, voice_over_text, backsound, durasi, saran_info) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)"
                for r in formatted_scenes:
                  cursor.execute(
                      query_scene,
                      (
                          sb_id,
                          r["Judul Scene"],
                          r["Visualisasi"],
                          r["Instruksi untuk Visual"],
                          r["Animasi"],
                          r["On-Screen Text"],
                          r["Voice Over Text"],
                          r["Backsound"],
                          r["Durasi"],
                          r["Saran dan Info"],
                      ),
                  )
                conn.commit()
                cursor.close()
                conn.close()
                st.success("✨ Berhasil digenerate dan direkam ke Railway MySQL!")
            else:
              st.error("Gagal mengurai format JSON dari LLM.")
          except Exception as e:
            st.error(f"Terjadi kesalahan saat pemrosesan LLM: {e}")

    if st.session_state.rag_contexts:
      for i, doc in enumerate(st.session_state.rag_contexts):
        st.markdown(
            f"""
                <div class="card-box">
                    <b>Referensi #{i+1}</b>
                    <p style="font-size: 13px;">{doc.page_content[:180]}...</p>
                </div>
            """,
            unsafe_allow_html=True,
        )

  with col_right:
    st.subheader("Storyboard Editor & Management")
    if "storyboard_df" not in st.session_state:
      st.session_state.storyboard_df = pd.DataFrame(columns=[
          "Judul Scene",
          "Visualisasi",
          "Instruksi untuk Visual",
          "Animasi",
          "On-Screen Text",
          "Voice Over Text",
          "Backsound",
          "Durasi",
          "Saran dan Info",
      ])

    updated_rows = []
    for idx, row in st.session_state.storyboard_df.iterrows():
      with st.expander(
          f"🎬 {row.get('Judul Scene', f'Scene {idx+1}')}", expanded=False
      ):
        c1, c2 = st.columns(2)
        with c1:
          nj = st.text_input(
              "Judul Scene", value=str(row["Judul Scene"]), key=f"jdl_{idx}"
          )
          nd = st.text_input(
              "Durasi", value=str(row["Durasi"]), key=f"dur_{idx}"
          )
          na = st.text_input(
              "Animasi", value=str(row["Animasi"]), key=f"anim_{idx}"
          )
          nost = st.text_input(
              "On-Screen Text",
              value=str(row["On-Screen Text"]),
              key=f"ost_{idx}",
          )
          nbs = st.text_input(
              "Backsound", value=str(row["Backsound"]), key=f"bs_{idx}"
          )
        with c2:
          nv = st.text_area(
              "Visualisasi", value=str(row["Visualisasi"]), key=f"vis_{idx}"
          )
          ni = st.text_area(
              "Instruksi Visual",
              value=str(row["Instruksi untuk Visual"]),
              key=f"inst_{idx}",
          )
          nvot = st.text_area(
              "Voice Over",
              value=str(row["Voice Over Text"]),
              key=f"vot_{idx}",
          )
          ns = st.text_area(
              "Saran", value=str(row["Saran dan Info"]), key=f"saran_{idx}"
          )

        updated_rows.append({
            "Judul Scene": nj,
            "Visualisasi": nv,
            "Instruksi untuk Visual": ni,
            "Animasi": na,
            "On-Screen Text": nost,
            "Voice Over Text": nvot,
            "Backsound": nbs,
            "Durasi": nd,
            "Saran dan Info": ns,
        })
    if updated_rows:
      st.session_state.storyboard_df = pd.DataFrame(updated_rows)

    if not st.session_state.storyboard_df.empty:
      pptx_data = generate_pptx(
          st.session_state.storyboard_df,
          program_studi,
          nama_mata_kuliah,
          project_name,
      )
      st.download_button(
          label="📥 Export to PPT",
          data=pptx_data,
          file_name="Storyboard.pptx",
          mime=(
              "application/vnd.openxmlformats-officedocument.presentationml.presentation"
          ),
          use_container_width=True,
      )
